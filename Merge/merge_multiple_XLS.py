import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import Font

# Lista dei file (possono essere xls o xlsx)
files = ['file1.xls', 'file2.xlsx', 'file3.xls']

# Attenzione al formato per stile e grandezza carattere (passo 2)
output_file = 'file_unificato.xlsx'

# 1 - Unisce i file convertendo .xls o .xlsx in dict con chiave il nome del foglio e valore il DataFrame
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
	for f in files:
		est = os.path.splitext(f)[1].lower()

		# Scegli il motore corretto in base all'estensione
		if est == '.xls':
			fogli = pd.read_excel(f, sheet_name=None, engine='xlrd')
		else:
			fogli = pd.read_excel(f, sheet_name=None)  # xlsx

		# Copia i fogli mantenendo i nomi (con suffisso per evitare conflitti)
		for nome_foglio, df in fogli.items():
			foglio_output = f'{nome_foglio}_{os.path.splitext(os.path.basename(f))[0]}'
			df.to_excel(writer, sheet_name=foglio_output, index=False)

# 2 - Applica stile, grandezza del carattere (vale solo per xlsx) e auto-fit delle colonne
wb = load_workbook(output_file)

# Scegli lo stile del font
font_style = Font(
	name="Courier New",   # Nome del font
	size=8,		  # Grandezza del carattere
	bold=False,	   # Grassetto
	italic=False,	 # Corsivo
	color="000000"	# Colore (esadecimale)
)

for sheet in wb.sheetnames:
	ws = wb[sheet]
	
	# Applica il font a tutte le celle di tutti i fogli
	for row in ws.iter_rows():
		for cell in row:
			cell.font = font_style
			
	# Applica auto-fit alle colonne 
	for col in ws.columns:
		max_length = 0
		col_letter = col[0].column_letter

		for cell in col:
			if cell.value:
				max_length = max(max_length, len(str(cell.value)))

		# Margine extra per leggibilità
		ws.column_dimensions[col_letter].width = max_length + 2

# Salva il file
wb.save(output_file)

print("File creato: " + output_file)